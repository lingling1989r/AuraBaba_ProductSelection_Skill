#!/usr/bin/env python3
"""Assemble every stage table into one Excel workbook.

Sheets
------
  00_说明          run metadata + how to read the workbook
  01_数据源能力      S0 source capability matrix
  02_热词发现        S1 seed keywords
  03_趋势时间序列     S2 raw monthly series
  04_趋势汇总        S2 per-keyword CAGR / seasonality / verdict
  05_场景关键词       S3 scene keywords
  06_子圈层          S3 chosen sub-community
  07_语料索引        S4 corpus index
  08_痛点聚类        S4 pain points
  09_市场概览        S5 market overview
  10_价格带          S5 price bands
  11_竞品            S5 competitors
  12_机会评分        S6 candidate scores
  13_阶段验证        gate verdicts for every stage
  14_溯源            per-file row counts + which tier each row came from
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=13, color="1F3864")

SHEETS: list[tuple[str, str]] = [
    ("01_数据源能力", "data/00_source_capability.csv"),
    ("02_热词发现", "data/01_seed_keywords.csv"),
    ("03_趋势时间序列", "data/02_trend_timeseries.csv"),
    ("04_趋势汇总", "data/02_trend_summary.csv"),
    ("05_场景关键词", "data/03_scene_keywords.csv"),
    ("06_子圈层", "data/03_subcommunity.csv"),
    ("07_语料索引", "data/04_corpus_index.csv"),
    ("08_痛点聚类", "data/04_painpoints.csv"),
    ("09_市场概览", "data/05_market_overview.csv"),
    ("10_价格带", "data/05_price_band.csv"),
    ("11_竞品", "data/05_competitors.csv"),
    ("12_机会评分", "data/06_candidate_scores.csv"),
]


def read_csv(path: Path) -> tuple[list[str], list[dict]]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        rdr = csv.DictReader(f)
        return (rdr.fieldnames or []), [r for r in rdr]


def autosize(ws, cols: list[str], rows: list[dict]) -> None:
    for idx, col in enumerate(cols, start=1):
        width = len(str(col)) + 4
        for r in rows[:200]:
            width = max(width, min(len(str(r.get(col, ""))) + 2, 60))
        ws.column_dimensions[get_column_letter(idx)].width = width


def add_table(ws, cols: list[str], rows: list[dict], title: str | None = None) -> int:
    r0 = 1
    if title:
        ws.cell(row=1, column=1, value=title).font = TITLE_FONT
        r0 = 3
    if not cols:
        ws.cell(row=r0, column=1, value="(empty - stage not produced)")
        return r0
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=r0, column=j, value=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center")
    for i, row in enumerate(rows, start=r0 + 1):
        for j, c in enumerate(cols, start=1):
            ws.cell(row=i, column=j, value=row.get(c, ""))
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    autosize(ws, cols, rows)
    return len(rows)


def build(run_dir: Path) -> Path:
    run_dir = run_dir.resolve()
    out_dir = run_dir / "out"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "selection_data.xlsx"

    state = json.loads((run_dir / "state.json").read_text(encoding="utf-8"))
    wb = Workbook()

    # ---- 00_说明
    ws = wb.active
    ws.title = "00_说明"
    ws.cell(row=1, column=1, value="社媒 x 关键词选品 - 全量数据工作簿").font = TITLE_FONT
    meta = [
        ("主题", state.get("topic", "")),
        ("目标市场", state.get("market", "")),
        ("类目", state.get("category", "")),
        ("起始关键词", ", ".join(state.get("seed_keywords") or [])),
        ("创建时间", state.get("created_at", "")),
        ("导出时间", state.get("updated_at", "")),
        ("", ""),
        ("数据原则", "先出数据再分析：每个 sheet 只承载某一阶段从数据源取回的原始/规范化表，结论在 HTML 报告中。"),
        ("降级链", "MCP(一级) -> Apify(二级) -> WebFetch/WebSearch(三级) -> 浏览器操作(四级)"),
        ("溯源", "每行都带 source_tier / captured_at，见 14_溯源 sheet。"),
    ]
    for i, (k, v) in enumerate(meta, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 100
    ws.cell(row=3 + len(meta) + 1, column=1, value="阶段产出概览").font = TITLE_FONT
    hdr_row = 3 + len(meta) + 3
    for j, h in enumerate(["阶段", "名称", "状态", "验证"], start=1):
        c = ws.cell(row=hdr_row, column=j, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    prov_rows = []
    for i, s in enumerate(state.get("stages", {}).values(), start=hdr_row + 1):
        ws.cell(row=i, column=1, value=s["id"])
        ws.cell(row=i, column=2, value=s["name"])
        ws.cell(row=i, column=3, value=s["status"])
        gate = s.get("gate") or {}
        ws.cell(row=i, column=4, value="PASS" if gate.get("ok") else ("FAIL" if gate else "-"))

    # ---- per-stage data sheets
    for sheet_name, rel in SHEETS:
        w = wb.create_sheet(sheet_name)
        p = run_dir / rel
        if p.is_file():
            cols, rows = read_csv(p)
            add_table(w, cols, rows, f"{sheet_name}  <-  {rel}")
            tiers: dict[str, int] = {}
            for r in rows:
                t = str(r.get("source_tier") or "-")
                tiers[t] = tiers.get(t, 0) + 1
            prov_rows.append({"file": rel, "rows": len(rows), "columns": len(cols),
                              "tiers": "; ".join(f"{k}={v}" for k, v in sorted(tiers.items()))})
        else:
            add_table(w, [], [], f"{sheet_name}  <-  {rel}")
            prov_rows.append({"file": rel, "rows": 0, "columns": 0, "tiers": "stage not produced"})

    # ---- 13_阶段验证
    w = wb.create_sheet("13_阶段验证")
    gr = []
    for s in state.get("stages", {}).values():
        gate = s.get("gate") or {}
        for c in gate.get("checks", []):
            gr.append({"stage": s["id"], "status": s["status"], "check": c.get("check", ""),
                       "ok": "PASS" if c.get("ok") else "FAIL", "path": c.get("path", ""),
                       "detail": c.get("detail", "")})
    add_table(w, ["stage", "status", "check", "ok", "path", "detail"], gr, "13_阶段验证 - 每阶段 gate 逐条结果")

    # ---- 14_溯源
    w = wb.create_sheet("14_溯源")
    add_table(w, ["file", "rows", "columns", "tiers"], prov_rows,
              "14_溯源 - 每个数据文件的行数与命中的数据源层级")

    wb.save(out_path)
    return out_path


def main() -> int:
    import argparse

    ap = argparse.ArgumentParser()
    ap.add_argument("--run-dir", default="./selection_run")
    args = ap.parse_args()
    p = build(Path(args.run_dir))
    print(f"wrote {p} ({p.stat().st_size:,} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
